"""
Export Jira sub-task worklogs into a flat Excel file for QA.

Each worklog entry becomes its own row (multiple rows per issue).
"""
from __future__ import annotations

import argparse
import os
import sqlite3
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from ipp_meeting_utils import (
    fetch_jira_issue_planned_dates,
    load_ipp_actual_and_remarks_by_key,
    load_ipp_issue_keys,
    load_ipp_planned_dates_by_key,
    yes_no_dates_altered,
    yes_no_ipp_actual_matches_jira_end,
    yes_no_in_ipp,
)
from jira_incremental_cache import (
    apply_overlap,
    bootstrap_default_checkpoint,
    get_cached_issue_payloads,
    get_cached_worklogs_for_subtasks,
    get_changed_or_new_issue_keys,
    get_db_path,
    get_or_init_checkpoint,
    init_db,
    mark_full_sync,
    mark_missing_issues_deleted,
    needs_full_sync,
    parse_iso_utc,
    record_pipeline_artifact,
    set_checkpoint,
    upsert_issue_index,
    upsert_issue_payloads,
    upsert_worklog_payload,
    utc_now_iso,
)
from jira_client import BASE_URL, get_session

DEFAULT_OUTPUT = "2_jira_subtask_worklogs.xlsx"
# Empty list means all projects.
DEFAULT_PROJECT_KEYS: list[str] = []
# Empty list means all issue types.
DEFAULT_WORKLOG_ISSUE_TYPES: list[str] = []


def _resolve_capacity_settings_db_path() -> Path:
    value = (os.getenv("JIRA_ASSIGNEE_HOURS_CAPACITY_DB_PATH", "assignee_hours_capacity.db") or "").strip()
    path = Path(value or "assignee_hours_capacity.db")
    if path.is_absolute():
        return path
    return Path(__file__).resolve().parent / path


def _load_project_keys_from_managed_db() -> list[str]:
    db_path = _resolve_capacity_settings_db_path()
    if not db_path.exists():
        return []
    conn = sqlite3.connect(db_path)
    try:
        rows = conn.execute(
            """
            SELECT project_key
            FROM managed_projects
            WHERE is_active = 1
            ORDER BY project_key ASC
            """
        ).fetchall()
    except sqlite3.Error:
        return []
    finally:
        conn.close()
    return [str(row[0]).strip().upper() for row in rows if str(row[0]).strip()]


def _get_project_keys() -> tuple[list[str], str]:
    db_keys = _load_project_keys_from_managed_db()
    if db_keys:
        return db_keys, "managed_projects_db"
    raw = os.getenv("JIRA_PROJECT_KEYS", "")
    if not raw.strip():
        return list(DEFAULT_PROJECT_KEYS), "env_fallback"
    return [key.strip() for key in raw.split(",") if key.strip()], "env_fallback"


def _get_worklog_issue_types() -> list[str]:
    raw = os.getenv("JIRA_WORKLOG_ISSUE_TYPES", "").strip()
    if not raw:
        return list(DEFAULT_WORKLOG_ISSUE_TYPES)
    parsed = [item.strip() for item in raw.split(",") if item.strip()]
    if len(parsed) == 1 and parsed[0].upper() in {"ALL", "*"}:
        return []
    return parsed


def _is_incremental_disabled(enable_incremental: bool) -> bool:
    if enable_incremental:
        return False
    return (os.getenv("JIRA_INCREMENTAL_DISABLE", "1").strip() or "1") == "1"


def _get_overlap_minutes() -> int:
    raw = os.getenv("JIRA_INCREMENTAL_OVERLAP_MINUTES", "5").strip() or "5"
    try:
        return max(int(raw), 0)
    except ValueError:
        return 5


def _get_force_full_days() -> int:
    raw = os.getenv("JIRA_FORCE_FULL_SYNC_DAYS", "7").strip() or "7"
    try:
        return max(int(raw), 1)
    except ValueError:
        return 7


def _get_bootstrap_days() -> int:
    raw = os.getenv("JIRA_INCREMENTAL_BOOTSTRAP_DAYS", "365").strip() or "365"
    try:
        return max(int(raw), 1)
    except ValueError:
        return 365


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Jira worklogs to Excel.")
    parser.add_argument(
        "--incremental",
        action="store_true",
        help="Enable smart incremental fetch (default: full fetch).",
    )
    return parser.parse_args()


def _target_assignee() -> str:
    return (os.getenv("JIRA_TARGET_ASSIGNEE", "") or "").strip()


def _jql_escape(value: str) -> str:
    return (value or "").replace("\\", "\\\\").replace('"', '\\"')


def _to_jql_updated_value(value: str) -> str:
    return parse_iso_utc(value).strftime("%Y-%m-%d %H:%M")


def _fetch_issues(session, jql: str, fields: list[str]) -> list[dict]:
    url = f"{BASE_URL}/rest/api/3/search/jql"
    issues: list[dict] = []
    next_page_token = None

    while True:
        payload = {"jql": jql, "maxResults": 100, "fields": fields}
        if next_page_token:
            payload["nextPageToken"] = next_page_token

        response = None
        for attempt in range(6):
            response = session.post(url, json=payload)
            if response.status_code == 429:
                retry_after = response.headers.get("Retry-After")
                try:
                    retry_after_seconds = float(retry_after) if retry_after else 0.0
                except ValueError:
                    retry_after_seconds = 0.0
                backoff_seconds = max((0.5 * (2**attempt)), retry_after_seconds, 0.5)
                time.sleep(backoff_seconds)
                continue
            response.raise_for_status()
            break
        if response is None:
            break
        if response.status_code == 429:
            response.raise_for_status()
        data = response.json()

        issues.extend(data.get("issues", []))
        next_page_token = data.get("nextPageToken")
        if not next_page_token:
            break

    return issues


def _fetch_issues_by_keys(session, issue_keys: list[str], fields: list[str]) -> list[dict]:
    if not issue_keys:
        return []
    results: list[dict] = []
    for offset in range(0, len(issue_keys), 500):
        chunk = issue_keys[offset : offset + 500]
        keys_clause = ", ".join(f'"{k}"' for k in chunk)
        jql = f"key in ({keys_clause})"
        results.extend(_fetch_issues(session, jql=jql, fields=fields))
    return results


def _build_discovery_jql(
    project_keys: list[str],
    issue_types: list[str],
    from_updated_utc: str | None,
    assignee_name: str = "",
) -> str:
    clauses: list[str] = []
    if project_keys:
        keys_str = ", ".join(project_keys)
        clauses.append(f"project in ({keys_str})")
    if issue_types:
        kinds = ", ".join(f'"{item}"' for item in issue_types)
        clauses.append(f"issuetype in ({kinds})")
    if assignee_name:
        clauses.append(f'assignee = "{_jql_escape(assignee_name)}"')
    base = " AND ".join(clauses) if clauses else "issuekey IS NOT EMPTY"
    if from_updated_utc:
        return f'{base} AND updated >= "{_to_jql_updated_value(from_updated_utc)}" ORDER BY updated ASC'
    return f"{base} ORDER BY updated ASC"


def _candidate_rows_from_issues(issues: list[dict]) -> list[dict]:
    now_utc = utc_now_iso()
    rows: list[dict] = []
    for issue in issues:
        fields = issue.get("fields", {}) or {}
        rows.append(
            {
                "issue_id": str(issue.get("id", "")).strip(),
                "issue_key": str(issue.get("key", "")).strip(),
                "updated_utc": str(fields.get("updated", "")).strip(),
                "issue_type": str((fields.get("issuetype") or {}).get("name", "")).strip(),
                "project_key": str((fields.get("project") or {}).get("key", "")).strip(),
                "last_seen_utc": now_utc,
                "is_deleted": 0,
            }
        )
    return [row for row in rows if row["issue_id"] and row["issue_key"]]


def _get_active_issue_keys(
    conn: sqlite3.Connection,
    project_keys: list[str],
    issue_types: list[str],
) -> list[str]:
    where = ["is_deleted = 0"]
    params: list[str] = []
    if project_keys:
        key_placeholders = ",".join("?" for _ in project_keys)
        where.append(f"project_key IN ({key_placeholders})")
        params.extend(project_keys)
    if issue_types:
        type_placeholders = ",".join("?" for _ in issue_types)
        where.append(f"issue_type IN ({type_placeholders})")
        params.extend(issue_types)
    rows = conn.execute(
        f"""
        SELECT issue_key
        FROM issue_index
        WHERE {' AND '.join(where)}
        ORDER BY issue_key
        """,
        tuple(params),
    ).fetchall()
    return [str(row[0]) for row in rows]


def _classify_candidates(
    conn: sqlite3.Connection,
    candidates: list[dict],
) -> tuple[list[str], int, int]:
    issue_ids = [str(item.get("issue_id", "")).strip() for item in candidates if str(item.get("issue_id", "")).strip()]
    existing_updated: dict[str, str] = {}
    for offset in range(0, len(issue_ids), 900):
        chunk = issue_ids[offset : offset + 900]
        if not chunk:
            continue
        placeholders = ",".join("?" for _ in chunk)
        rows = conn.execute(
            f"SELECT issue_id, updated_utc FROM issue_index WHERE issue_id IN ({placeholders})",
            tuple(chunk),
        ).fetchall()
        for row in rows:
            existing_updated[str(row[0])] = str(row[1] or "")

    changed: list[str] = []
    seen: set[str] = set()
    new_count = 0
    changed_existing_count = 0
    for item in candidates:
        issue_id = str(item.get("issue_id", "")).strip()
        issue_key = str(item.get("issue_key", "")).strip()
        updated_utc = str(item.get("updated_utc", "")).strip()
        if not issue_id or not issue_key:
            continue
        if issue_id not in existing_updated:
            new_count += 1
            if issue_key not in seen:
                changed.append(issue_key)
                seen.add(issue_key)
            continue
        if existing_updated.get(issue_id, "") != updated_utc:
            changed_existing_count += 1
            if issue_key not in seen:
                changed.append(issue_key)
                seen.add(issue_key)
    return changed, new_count, changed_existing_count


def _get_epic_key_from_issue_fields(fields: dict) -> str:
    parent = fields.get("parent") or {}
    parent_key = parent.get("key")
    if parent_key:
        return parent_key

    epic_link = fields.get("customfield_10014")
    if isinstance(epic_link, str) and epic_link.strip():
        return epic_link.strip()
    if isinstance(epic_link, dict):
        epic_key = epic_link.get("key")
        if epic_key:
            return epic_key
    return ""


def _fetch_all_worklogs_for_issue(
    session,
    issue_key: str,
    delay_seconds: float = 0.2,
    max_retries: int = 5,
    request_timeout_seconds: float = 30.0,
) -> list[dict]:
    url = f"{BASE_URL}/rest/api/3/issue/{issue_key}/worklog"
    start_at = 0
    max_results = 100
    all_logs: list[dict] = []

    while True:
        params = {"startAt": start_at, "maxResults": max_results}
        response = None
        for attempt in range(max_retries + 1):
            try:
                response = session.get(url, params=params, timeout=request_timeout_seconds)
            except Exception as e:
                if attempt >= max_retries:
                    raise
                backoff_seconds = max(delay_seconds * (2 ** attempt), 0.5)
                print(
                    f"Worklog request timeout/error for {issue_key} (attempt {attempt + 1}/{max_retries + 1}): {e}; "
                    f"sleeping {backoff_seconds:.2f}s"
                )
                time.sleep(backoff_seconds)
                continue
            if response.status_code == 429:
                if attempt >= max_retries:
                    response.raise_for_status()
                retry_after = response.headers.get("Retry-After")
                try:
                    retry_after_seconds = float(retry_after) if retry_after else 0.0
                except ValueError:
                    retry_after_seconds = 0.0
                backoff_seconds = max(delay_seconds * (2 ** attempt), retry_after_seconds, 0.5)
                print(
                    f"Rate limited for {issue_key} (attempt {attempt + 1}/{max_retries + 1}); "
                    f"sleeping {backoff_seconds:.2f}s"
                )
                time.sleep(backoff_seconds)
                continue
            response.raise_for_status()
            break

        if response is None:
            break

        data = response.json()

        page_logs = data.get("worklogs", []) or []
        all_logs.extend(page_logs)

        total = data.get("total", len(all_logs))
        start_at += len(page_logs)
        if start_at >= total or not page_logs:
            break

        if delay_seconds > 0:
            time.sleep(delay_seconds)

    return all_logs


def _seconds_to_hours(seconds_value) -> float:
    if seconds_value in (None, ""):
        return 0.0
    try:
        return round(float(seconds_value) / 3600.0, 2)
    except (TypeError, ValueError):
        return 0.0


def _write_excel(rows: list[list[object]], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SubtaskWorklogs"

    headers = [
        "issue_link",
        "issue_id",
        "issue_title",
        "issue_type",
        "parent_story_link",
        "parent_story_id",
        "parent_epic_id",
        "issue_assignee",
        "Latest IPP Meeting",
        "Jira IPP RMI Dates Altered",
        "IPP Actual Date (Production Date)",
        "IPP Remarks",
        "IPP Actual Date Matches Jira End Date",
        "worklog_started",
        "hours_logged",
        "worklog_author",
    ]
    sheet.append(headers)

    for row in rows:
        sheet.append(row)

    workbook.save(output_path)


def main() -> None:
    args = _parse_args()
    run_started = datetime.now(timezone.utc)
    run_id = uuid.uuid4().hex
    target_assignee = _target_assignee()
    session = get_session()
    project_keys, project_source = _get_project_keys()
    issue_types = _get_worklog_issue_types()
    incremental_disabled = _is_incremental_disabled(args.incremental)
    overlap_minutes = _get_overlap_minutes()
    force_full_days = _get_force_full_days()
    bootstrap_days = _get_bootstrap_days()
    db_path = get_db_path()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    init_db(conn)

    default_checkpoint = bootstrap_default_checkpoint(bootstrap_days)
    checkpoint = get_or_init_checkpoint(conn, "worklogs", default_checkpoint)
    now_utc = utc_now_iso()
    force_full_sync = incremental_disabled or needs_full_sync(conn, "worklogs", now_utc, force_full_days)
    from_updated = None if force_full_sync else apply_overlap(checkpoint, overlap_minutes)

    print(f"Incremental mode: {'OFF' if incremental_disabled else 'ON'}")
    print(f"Sync DB: {db_path}")
    print(f"Projects source: {project_source}")
    print(f"Force full sync: {'Yes' if force_full_sync else 'No'}")
    print(f"Projects: {', '.join(project_keys) if project_keys else 'ALL'}")
    print(f"Issue types: {', '.join(issue_types) if issue_types else 'ALL'}")
    if from_updated:
        print(f"Discovery updated >= {from_updated} (checkpoint={checkpoint}, overlap={overlap_minutes}m)")

    discovery_fields = [
        "project",
        "issuetype",
        "updated",
    ]
    discovery_jql = _build_discovery_jql(
        project_keys,
        issue_types,
        from_updated_utc=from_updated,
        assignee_name=target_assignee,
    )
    print(f"Running discovery query")
    discovered = _fetch_issues(session, jql=discovery_jql, fields=discovery_fields)
    candidates = _candidate_rows_from_issues(discovered)
    changed_issue_keys, new_issue_count, changed_existing_count = _classify_candidates(conn, candidates)
    upsert_issue_index(conn, candidates)

    active_ids = {row["issue_id"] for row in candidates}
    deleted_count = 0
    if force_full_sync and not target_assignee:
        deleted_count = mark_missing_issues_deleted(conn, project_keys, issue_types, active_ids)
        if deleted_count:
            print(f"Marked deleted/inaccessible issues: {deleted_count}")

    # Keep API-compatible classification path available as guard in case of classification drift.
    if not force_full_sync and not changed_issue_keys:
        changed_issue_keys = get_changed_or_new_issue_keys(conn, candidates)
    discovered_keys = [str(item.get("issue_key", "")).strip() for item in candidates if str(item.get("issue_key", "")).strip()]
    detail_fetch_keys = discovered_keys if force_full_sync else changed_issue_keys
    detail_fetch_keys = sorted({key for key in detail_fetch_keys if key})

    detail_fields = [
        "project",
        "summary",
        "issuetype",
        "parent",
        "assignee",
        "timespent",
        "updated",
    ]
    if detail_fetch_keys:
        print(f"Fetching full issue payloads for {len(detail_fetch_keys)} changed/new issues")
        detailed = _fetch_issues_by_keys(session, detail_fetch_keys, detail_fields)
        payload_rows: list[dict] = []
        index_rows: list[dict] = []
        now_seen = utc_now_iso()
        for issue in detailed:
            issue_id = str(issue.get("id", "")).strip()
            issue_key = str(issue.get("key", "")).strip()
            fields = issue.get("fields", {}) or {}
            updated_utc = str(fields.get("updated", "")).strip()
            issue_type = str((fields.get("issuetype") or {}).get("name", "")).strip()
            project_key = str((fields.get("project") or {}).get("key", "")).strip()
            if not issue_id or not issue_key:
                continue
            payload_rows.append(
                {
                    "issue_id": issue_id,
                    "issue_key": issue_key,
                    "updated_utc": updated_utc,
                    "payload": issue,
                }
            )
            index_rows.append(
                {
                    "issue_id": issue_id,
                    "issue_key": issue_key,
                    "updated_utc": updated_utc,
                    "issue_type": issue_type,
                    "project_key": project_key,
                    "last_seen_utc": now_seen,
                    "is_deleted": 0,
                }
            )
        upsert_issue_payloads(conn, payload_rows)
        upsert_issue_index(conn, index_rows)

    active_subtask_keys = discovered_keys if target_assignee else _get_active_issue_keys(conn, project_keys, issue_types)
    cached_payloads = get_cached_issue_payloads(conn, project_keys=project_keys, issue_types=issue_types)
    if target_assignee:
        target_assignee_lower = target_assignee.casefold()
        cached_payloads = [
            item
            for item in cached_payloads
            if str((((item.get("fields") or {}).get("assignee") or {}).get("displayName") or "")).strip().casefold() == target_assignee_lower
        ]
    cached_payload_by_key = {str(item.get("key", "")).strip(): item for item in cached_payloads if str(item.get("key", "")).strip()}
    missing_payload_keys = [key for key in active_subtask_keys if key not in cached_payload_by_key]
    if missing_payload_keys:
        print(f"Backfilling missing cached payloads for {len(missing_payload_keys)} issues")
        missing_detailed = _fetch_issues_by_keys(session, sorted(missing_payload_keys), detail_fields)
        payload_rows = []
        index_rows = []
        now_seen = utc_now_iso()
        for issue in missing_detailed:
            issue_id = str(issue.get("id", "")).strip()
            issue_key = str(issue.get("key", "")).strip()
            fields = issue.get("fields", {}) or {}
            updated_utc = str(fields.get("updated", "")).strip()
            issue_type = str((fields.get("issuetype") or {}).get("name", "")).strip()
            project_key = str((fields.get("project") or {}).get("key", "")).strip()
            if not issue_id or not issue_key:
                continue
            payload_rows.append(
                {
                    "issue_id": issue_id,
                    "issue_key": issue_key,
                    "updated_utc": updated_utc,
                    "payload": issue,
                }
            )
            index_rows.append(
                {
                    "issue_id": issue_id,
                    "issue_key": issue_key,
                    "updated_utc": updated_utc,
                    "issue_type": issue_type,
                    "project_key": project_key,
                    "last_seen_utc": now_seen,
                    "is_deleted": 0,
                }
            )
        upsert_issue_payloads(conn, payload_rows)
        upsert_issue_index(conn, index_rows)

    subtasks = get_cached_issue_payloads(conn, project_keys=project_keys, issue_types=issue_types)
    if target_assignee:
        target_assignee_lower = target_assignee.casefold()
        subtasks = [
            item
            for item in subtasks
            if str((((item.get("fields") or {}).get("assignee") or {}).get("displayName") or "")).strip().casefold() == target_assignee_lower
        ]
    print(f"Issue payloads available in cache: {len(subtasks)}")
    parent_story_keys: set[str] = set()
    subtask_info: dict[str, dict] = {}
    for issue in subtasks:
        issue_key = issue.get("key", "")
        issue_fields = issue.get("fields", {}) or {}
        parent_story_key = (issue_fields.get("parent") or {}).get("key", "") or ""
        if parent_story_key:
            parent_story_keys.add(parent_story_key)
        subtask_info[issue_key] = {
            "summary": issue_fields.get("summary", "") or "",
            "issue_type": (issue_fields.get("issuetype") or {}).get("name", "") or "",
            "assignee": (issue_fields.get("assignee") or {}).get("displayName", "Unassigned"),
            "parent_story_key": parent_story_key,
            "timespent_seconds": issue_fields.get("timespent") or 0,
            "issue_id": str(issue.get("id", "")).strip(),
        }

    story_to_epic: dict[str, str] = {}
    if parent_story_keys:
        print(f"Fetching parent story details for {len(parent_story_keys)} stories")
        story_fields = ["parent", "customfield_10014", "issuetype"]
        batch = list(parent_story_keys)
        for i in range(0, len(batch), 500):
            chunk = batch[i : i + 500]
            keys = ", ".join(f'"{k}"' for k in chunk)
            story_jql = f"key in ({keys})"
            stories = _fetch_issues(session, jql=story_jql, fields=story_fields)
            for story in stories:
                story_key = story.get("key", "")
                fields_dict = story.get("fields", {}) or {}
                story_to_epic[story_key] = _get_epic_key_from_issue_fields(fields_dict)

    rows: list[list[object]] = []
    max_issues_env = os.getenv("JIRA_WORKLOG_MAX_ISSUES", "").strip()
    max_issues = int(max_issues_env) if max_issues_env.isdigit() else None
    delay_seconds_env = os.getenv("JIRA_WORKLOG_DELAY_SECONDS", "0.2").strip() or "0.2"
    max_retries_env = os.getenv("JIRA_WORKLOG_MAX_RETRIES", "5").strip() or "5"
    request_timeout_env = os.getenv("JIRA_WORKLOG_REQUEST_TIMEOUT_SECONDS", "30").strip() or "30"
    try:
        delay_seconds = max(float(delay_seconds_env), 0.0)
    except ValueError:
        delay_seconds = 0.2
    try:
        max_retries = max(int(max_retries_env), 0)
    except ValueError:
        max_retries = 5
    try:
        request_timeout_seconds = max(float(request_timeout_env), 1.0)
    except ValueError:
        request_timeout_seconds = 30.0
    skipped_no_worklogs = 0

    ipp_issue_keys = load_ipp_issue_keys()
    print(f"IPP issue keys loaded: {len(ipp_issue_keys)}")
    ipp_planned_dates = load_ipp_planned_dates_by_key()
    print(f"IPP planned-date keys loaded: {len(ipp_planned_dates)}")
    ipp_actual_by_key = load_ipp_actual_and_remarks_by_key()
    print(f"IPP actual-date/remarks keys loaded: {len(ipp_actual_by_key)}")
    jira_epic_keys = {key for key in story_to_epic.values() if key}
    jira_epic_dates = fetch_jira_issue_planned_dates(session, BASE_URL, jira_epic_keys)
    print(f"Jira epic planned-date keys loaded: {len(jira_epic_dates)}")

    all_subtask_keys = sorted(subtask_info.keys())
    cached_worklogs = get_cached_worklogs_for_subtasks(conn, all_subtask_keys)
    worklog_fetch_keys = set(detail_fetch_keys if not force_full_sync else all_subtask_keys)
    worklog_fetch_keys.update(key for key in all_subtask_keys if key not in cached_worklogs)
    worklog_fetch_count = 0
    processed = 0
    for issue_key in all_subtask_keys:
        info = subtask_info[issue_key]
        processed += 1
        if max_issues is not None and processed > max_issues:
            break
        if processed % 25 == 0:
            print(f"Worklog fetch progress: issue {processed}/{len(subtask_info)} ({issue_key})")

        issue_link = f"{BASE_URL}/browse/{issue_key}"
        parent_story_key = info.get("parent_story_key", "") or ""
        parent_story_link = f"{BASE_URL}/browse/{parent_story_key}" if parent_story_key else ""
        parent_epic_id = story_to_epic.get(parent_story_key, "") if parent_story_key else ""
        parent_epic_key = str(parent_epic_id or "").strip().upper()
        ipp_actual_data = ipp_actual_by_key.get(parent_epic_key, {})
        timespent_seconds = int(info.get("timespent_seconds") or 0)

        if timespent_seconds <= 0:
            skipped_no_worklogs += 1
            if issue_key in worklog_fetch_keys:
                upsert_worklog_payload(
                    conn,
                    issue_key=issue_key,
                    issue_id=str(info.get("issue_id", "")).strip(),
                    worklogs=[],
                    worklog_updated_utc=None,
                )
                cached_worklogs[issue_key] = []
            continue

        if issue_key in worklog_fetch_keys:
            try:
                worklogs = _fetch_all_worklogs_for_issue(
                    session,
                    issue_key,
                    delay_seconds=delay_seconds,
                    max_retries=max_retries,
                    request_timeout_seconds=request_timeout_seconds,
                )
                worklog_fetch_count += 1
                max_worklog_updated = ""
                for wl in worklogs:
                    wl_updated = str(wl.get("updated", "")).strip()
                    if wl_updated and (not max_worklog_updated or parse_iso_utc(wl_updated) > parse_iso_utc(max_worklog_updated)):
                        max_worklog_updated = wl_updated
                upsert_worklog_payload(
                    conn,
                    issue_key=issue_key,
                    issue_id=str(info.get("issue_id", "")).strip(),
                    worklogs=worklogs,
                    worklog_updated_utc=max_worklog_updated or None,
                )
                cached_worklogs[issue_key] = worklogs
            except Exception as e:
                print(f"Warning: could not fetch worklogs for {issue_key}: {e}")
                worklogs = cached_worklogs.get(issue_key, [])
        else:
            worklogs = cached_worklogs.get(issue_key, [])

        for wl in worklogs:
            started = wl.get("started", "") or ""
            hours = _seconds_to_hours(wl.get("timeSpentSeconds"))
            author = wl.get("author", {}) or {}
            worklog_author = (
                str(author.get("displayName", "")).strip()
                or str(author.get("emailAddress", "")).strip()
                or str(author.get("accountId", "")).strip()
                or "Unknown"
            )
            rows.append(
                [
                    issue_link,
                    issue_key,
                    info.get("summary", ""),
                    info.get("issue_type", ""),
                    parent_story_link,
                    parent_story_key,
                    parent_epic_id,
                    info.get("assignee", "Unassigned"),
                    yes_no_in_ipp(issue_key, ipp_issue_keys),
                    yes_no_dates_altered(parent_epic_id, ipp_planned_dates, jira_epic_dates),
                    ipp_actual_data.get("ipp_actual_date", ""),
                    ipp_actual_data.get("ipp_remarks", ""),
                    yes_no_ipp_actual_matches_jira_end(parent_epic_key, ipp_actual_by_key, jira_epic_dates),
                    started,
                    hours,
                    worklog_author,
                ]
            )

        if processed % 50 == 0:
            print(f"Processed {processed}/{len(subtask_info)} issues (rows so far: {len(rows)})")

    print(
        "Worklog API tuning: "
        f"delay={delay_seconds}s, max_retries={max_retries}, skipped_no_worklogs={skipped_no_worklogs}, "
        f"worklog_fetches={worklog_fetch_count}"
    )

    max_updated_seen = ""
    for row in candidates:
        updated_utc = str(row.get("updated_utc", "")).strip()
        if not updated_utc:
            continue
        if not max_updated_seen or parse_iso_utc(updated_utc) > parse_iso_utc(max_updated_seen):
            max_updated_seen = updated_utc
    if max_updated_seen:
        set_checkpoint(conn, "worklogs", max_updated_seen)
    if force_full_sync:
        mark_full_sync(conn, "worklogs", utc_now_iso())

    output_name = os.getenv("JIRA_WORKLOG_XLSX_PATH", DEFAULT_OUTPUT).strip() or DEFAULT_OUTPUT
    output_path = Path(output_name)
    if not output_path.is_absolute():
        output_path = Path(__file__).resolve().parent / output_path

    _write_excel(rows, output_path)
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    print(f"Export written: {output_path}")
    print(f"Generated at: {generated_at}")
    print(f"Rows: {len(rows)}")

    run_ended = datetime.now(timezone.utc)
    record_pipeline_artifact(
        conn,
        {
            "run_id": run_id,
            "pipeline": "worklogs",
            "started_at_utc": run_started.replace(microsecond=0).isoformat().replace("+00:00", "Z"),
            "ended_at_utc": run_ended.replace(microsecond=0).isoformat().replace("+00:00", "Z"),
            "issues_scanned": len(candidates),
            "issues_changed": changed_existing_count,
            "new_issues": new_issue_count,
            "detail_fetches": len(detail_fetch_keys),
            "worklog_fetches": worklog_fetch_count,
            "duration_ms": int((run_ended - run_started).total_seconds() * 1000),
        },
    )
    conn.close()


if __name__ == "__main__":
    main()
